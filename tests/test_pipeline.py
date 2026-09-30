"""Automated checks that the pipeline's outputs are correct and agree with each other.

Run the pipeline first (python run_all.py), then from the repo root:
    python -m unittest discover -s tests -v

The key idea: the same total (revenue, gross profit, order count) is computed four different ways
(clean CSV -> SQLite -> Excel formulas -> web dashboard JSON). If they ever disagree, a test fails.
"""
import contextlib
import json
import re
import sqlite3
import sys
import unittest
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from paths import CLEAN_DIR, DB_PATH, DOCS_DIR, EXCEL_PATH, SENSITIVITY_DIR, SQL_RESULTS_DIR  # noqa: E402

CHANNELS = {"Storefront", "Night Market", "Vendor Event"}


class CleanDataTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.orders = pd.read_csv(CLEAN_DIR / "orders.csv")
        cls.items = pd.read_csv(CLEAN_DIR / "order_items.csv")
        cls.customers = pd.read_csv(CLEAN_DIR / "customers.csv")

    def test_order_ids_unique(self):
        self.assertTrue(self.orders.order_id.is_unique)

    def test_no_duplicate_lines_or_negative_quantities(self):
        self.assertFalse(self.items.duplicated().any())
        self.assertTrue((self.items.quantity > 0).all())

    def test_categories_standardized(self):
        self.assertTrue(set(self.orders.channel) <= CHANNELS)
        self.assertTrue(set(self.orders.payment_method) <= {"Card", "Cash", "Apple Pay", "Online"})
        self.assertTrue(set(self.orders.order_type) <= {"Dine-in", "Takeout", "Delivery App"})

    def test_money_fields_are_consistent(self):
        o = self.orders
        self.assertTrue(((o.subtotal - o.discount_amount - o.revenue).abs() < 0.011).all())
        self.assertTrue(((o.revenue - o.food_cost_total - o.gross_profit).abs() < 0.011).all())
        self.assertTrue((o.revenue > 0).all())

    def test_line_profit_ties_to_order_profit(self):
        self.assertAlmostEqual(self.items.line_profit.sum(), self.orders.gross_profit.sum(), delta=2.0)

    def test_referential_integrity(self):
        self.assertTrue(self.items.order_id.isin(self.orders.order_id).all())
        self.assertTrue(self.orders.customer_id.dropna().isin(self.customers.customer_id).all())

    def test_prep_times_are_plausible(self):
        p = self.orders.prep_time_minutes.dropna()
        self.assertTrue(((p > 0) & (p <= 120)).all())


class CrossSystemTotalsTests(unittest.TestCase):
    """Revenue, profit and order counts must match across CSV, SQLite, Excel and the web dashboard."""

    @classmethod
    def setUpClass(cls):
        o = pd.read_csv(CLEAN_DIR / "orders.csv")
        cls.revenue, cls.profit, cls.n = o.revenue.sum(), o.gross_profit.sum(), len(o)

    def test_sqlite_matches_csv(self):
        with contextlib.closing(sqlite3.connect(DB_PATH)) as c:
            rev, gp, n = c.execute("SELECT SUM(revenue), SUM(gross_profit), COUNT(*) FROM orders").fetchone()
        self.assertAlmostEqual(rev, self.revenue, places=2)
        self.assertAlmostEqual(gp, self.profit, places=2)
        self.assertEqual(n, self.n)

    def test_web_dashboard_matches_csv(self):
        html = (DOCS_DIR / "index.html").read_text()
        data = json.loads(re.search(r"const D = (\{.*?\});\n", html, re.S).group(1))
        self.assertAlmostEqual(sum(r[3] for r in data["mc"]), self.revenue, delta=1.0)
        self.assertAlmostEqual(sum(r[4] for r in data["mc"]), self.profit, delta=1.0)
        self.assertEqual(sum(r[2] for r in data["mc"]), self.n)

    def test_excel_cached_kpis_match_csv(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed (pip install -r requirements-dev.txt)")
        ws = openpyxl.load_workbook(EXCEL_PATH, data_only=True)["Analysis"]
        found = {ws.cell(r, 2).value: ws.cell(r, 3).value for r in range(1, 20)}
        self.assertAlmostEqual(found["Total Revenue"], self.revenue, places=2)
        self.assertAlmostEqual(found["Gross Profit"], self.profit, places=2)
        self.assertEqual(found["Total Orders"], self.n)


class OutputFilesTests(unittest.TestCase):
    def test_all_sql_results_exist_and_are_not_empty(self):
        files = sorted(SQL_RESULTS_DIR.glob("Q*.csv"))
        self.assertEqual(len(files), 15)
        for f in files:
            self.assertGreater(len(pd.read_csv(f)), 0, f.name)

    def test_event_roi_roughly_matches_between_sql_and_sensitivity(self):
        q3 = pd.read_csv(next(SQL_RESULTS_DIR.glob("Q03_*.csv")))
        summary = json.loads((SENSITIVITY_DIR / "sensitivity_summary.json").read_text())
        gp = {e["event_name"]: e["gp"] for e in summary["events"]}
        for r in q3.itertuples():
            self.assertAlmostEqual(r.gross_profit, gp[r.event_name], delta=1.0)


class SensitivityTests(unittest.TestCase):
    def test_labor_sensitivity_is_monotonic(self):
        df = pd.read_csv(SENSITIVITY_DIR / "labor_rate_sensitivity.csv")
        self.assertTrue(df.tue_wed_savings_one_fewer_staff.is_monotonic_increasing)
        self.assertTrue(df.storefront_contribution.is_monotonic_decreasing)

    def test_weekday_confidence_intervals_contain_the_mean(self):
        df = pd.read_csv(SENSITIVITY_DIR / "weekday_breakeven.csv")
        self.assertTrue(((df.ci95_low <= df.avg_gross_profit_per_day) & (df.avg_gross_profit_per_day <= df.ci95_high)).all())

    def test_scenarios_are_ordered_worst_to_best(self):
        df = pd.read_csv(SENSITIVITY_DIR / "scenario_2026.csv")
        money = df[~df.lever.str.startswith("Improvement")]
        self.assertTrue((money.worst_case <= money.typical_case).all())
        self.assertTrue((money.typical_case <= money.best_case).all())


if __name__ == "__main__":
    unittest.main()
